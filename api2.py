from openpyxl import load_workbook
import statistics

# Cargar la hoja de datos
def load_data(filepath):
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    data = rows[1:]
    return headers, data


# Filtrar registros
def filter_data(headers, data, departamento, municipio, cultivo, n):
    idx_dep = headers.index("Departamento")
    idx_mun = headers.index("Municipio")
    idx_cult = headers.index("Cultivo")

    filtered = [
        row for row in data
        if str(row[idx_dep]).strip().lower() == departamento.strip().lower()
        and str(row[idx_mun]).strip().lower() == municipio.strip().lower()
        and str(row[idx_cult]).strip().lower() == cultivo.strip().lower()
    ]

    return filtered[:n]


# Calcular medianas del cultivo
def calculate_medians(headers, filtered_data):

    idx_ph   = headers.index("pH agua:suelo 2,5:1,0")
    idx_p    = headers.index("Fósforo (P) Bray II mg/kg")
    idx_k    = headers.index("Potasio (K) intercambiable cmol(+)/kg")

    col_ph = [row[idx_ph] for row in filtered_data if isinstance(row[idx_ph], (int, float))]
    col_p  = [row[idx_p]  for row in filtered_data if isinstance(row[idx_p], (int, float))]
    col_k  = [row[idx_k]  for row in filtered_data if isinstance(row[idx_k], (int, float))]

    return {
        "mediana_ph": statistics.median(col_ph) if col_ph else None,
        "mediana_p":  statistics.median(col_p)  if col_p else None,
        "mediana_k":  statistics.median(col_k)  if col_k else None
    }


# Construir tabla final para UI
def build_output_table(headers, filtered_data, medians):

    idx_dep = headers.index("Departamento")
    idx_mun = headers.index("Municipio")
    idx_cult = headers.index("Cultivo")
    idx_topo = headers.index("Topografia")

    table = []
    for row in filtered_data:
        table.append([
            row[idx_dep],
            row[idx_mun],
            row[idx_cult],
            row[idx_topo],
        ])

    return table, medians

