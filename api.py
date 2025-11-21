import openpyxl
import numpy as np

excel = openpyxl.load_workbook("resultado_laboratorio_suelo.xlsx", read_only = True)

hoja = excel[excel.sheetnames[0]]
hoja.calculate_dimension(force=True)

propiedades_columnas = hoja.iter_rows(min_row=1,max_row=1,values_only=True)
propiedades = []

for propiedad in propiedades_columnas:
    for cell in propiedad:
        propiedades.append(cell)

propiedades_suelo = {propiedad:np.array([]) for propiedad in propiedades}


def consultar_registros(municipio, departamento, cultivo, numero_registros):
    breakpoint()
    filas = hoja.iter_rows(min_row=2,max_row=hoja.max_row,values_only=True)

    for fila in filas:
        for indice, propiedad in enumerate(propiedades_suelo.keys(), start = 4):
            print(f"{propiedad} - {propiedades_suelo.get(propiedades[indice])} - {fila[indice]})")

consultar_registros(1,1,1,1)

   
